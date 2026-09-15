#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Exportación Excel de la lista de clientes (página actual)."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

HEADER_FONT = Font(bold=True)

COLUMNS = (
    'ID',
    'Categoría',
    'Nombre',
    'Teléfono',
    'Dirección',
    'Barrio',
    'Ciudad',
    'IP Router',
    'IP Antea',
    'Incidencias',
)

COLUMN_WIDTHS = {
    'A': 8,
    'B': 14,
    'C': 28,
    'D': 16,
    'E': 32,
    'F': 18,
    'G': 16,
    'H': 16,
    'I': 16,
    'J': 12,
}


def _client_row(client):
    from core.app import categoria_cliente_label

    ciudad = ''
    if getattr(client, 'ciudad_row', None) is not None:
        ciudad = client.ciudad_row.nombre or ''
    return (
        client.id,
        categoria_cliente_label(client.categoria),
        client.nom or '',
        client.telephone or '',
        client.adresse or '',
        client.ville or '',
        ciudad,
        client.ip_router or '',
        client.ip_antea or '',
        len(client.incidents or []),
    )


def build_clients_list_workbook(clients):
    """Genera un workbook Excel con solo la tabla de clientes."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Clientes'

    for col_idx, header in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=header).font = HEADER_FONT

    for row_idx, client in enumerate(clients, start=2):
        for col_idx, value in enumerate(_client_row(client), start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = 'A2'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_clients_export_filename():
    """Nombre de archivo: clientes_{YYYYMMDD_HHMM}.xlsx"""
    return f"clientes_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
